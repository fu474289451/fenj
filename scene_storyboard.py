"""
Video Scene Storyboard Tool
视频镜头分析与分镜表生成工具
"""

import os
import re
import sys
import shutil
import tempfile
import threading
import platform
import subprocess
from pathlib import Path
from collections import namedtuple

import tkinter as tk
from tkinter import ttk, messagebox

import yt_dlp
import cv2
from scenedetect import open_video, SceneManager
from scenedetect.detectors import ContentDetector
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Font, Alignment, PatternFill
from PIL import Image as PILImage

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------
FRAME_OFFSET_SEC = 0.5
THUMBNAIL_WIDTH_PX = 160
DEFAULT_THRESHOLD = 27.0
SHOT_DIR_NAME = "shots"
EXCEL_FILENAME = "storyboard.xlsx"

SceneCut = namedtuple("SceneCut", ["index", "start_time", "end_time", "start_frame", "end_frame"])


# ---------------------------------------------------------------------------
# Utility functions
# ---------------------------------------------------------------------------

def sanitize_filename(name: str) -> str:
    """Remove characters not allowed in Windows filenames."""
    name = re.sub(r'[<>:"/\\|?*]', '_', name)
    name = name.strip('. ')
    if len(name) > 100:
        name = name[:100]
    return name or "untitled"


def make_phase_cb(progress_cb, phase_start: float, phase_end: float):
    """Map a local 0-100 progress to a global progress range."""
    def cb(msg: str, local_pct: float):
        global_pct = phase_start + (local_pct / 100.0) * (phase_end - phase_start)
        progress_cb(msg, global_pct)
    return cb


def format_timecode(seconds: float) -> str:
    """Format seconds as HH:MM:SS.mmm."""
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = seconds % 60
    return f"{h:02d}:{m:02d}:{s:06.3f}"


def get_output_dir(title: str) -> Path:
    """Create output directory on Desktop, with suffix if already exists."""
    desktop = Path.home() / "Desktop"
    if not desktop.exists():
        desktop = Path.home()

    base_name = sanitize_filename(title)
    output_dir = desktop / base_name
    counter = 1
    while output_dir.exists():
        output_dir = desktop / f"{base_name}_{counter}"
        counter += 1
    output_dir.mkdir(parents=True)
    return output_dir


def open_folder(path: str):
    """Open a folder in the system file manager."""
    if platform.system() == "Windows":
        os.startfile(path)
    elif platform.system() == "Darwin":
        subprocess.Popen(["open", path])
    else:
        subprocess.Popen(["xdg-open", path])


# ---------------------------------------------------------------------------
# Pipeline functions
# ---------------------------------------------------------------------------

def download_video(url: str, temp_dir: str, progress_cb) -> tuple:
    """Download video using yt-dlp. Returns (video_path, title)."""

    def ydl_hook(d):
        if d["status"] == "downloading":
            total = d.get("total_bytes") or d.get("total_bytes_estimate") or 0
            downloaded = d.get("downloaded_bytes", 0)
            if total > 0:
                progress_cb("正在下载视频...", (downloaded / total) * 100)
        elif d["status"] == "finished":
            progress_cb("下载完成，正在处理...", 100)

    ydl_opts = {
        "outtmpl": os.path.join(temp_dir, "%(title)s.%(ext)s"),
        "format": "bestvideo[ext=mp4]+bestaudio[ext=m4a]/best[ext=mp4]/best",
        "merge_output_format": "mp4",
        "progress_hooks": [ydl_hook],
        "noplaylist": True,
        "quiet": True,
        "no_warnings": True,
    }

    try:
        with yt_dlp.YoutubeDL(ydl_opts) as ydl:
            info = ydl.extract_info(url, download=True)
            title = info.get("title", "untitled")
            filename = ydl.prepare_filename(info)
            # yt-dlp may merge to .mp4
            video_path = os.path.splitext(filename)[0] + ".mp4"
            if not os.path.exists(video_path):
                video_path = filename
            if not os.path.exists(video_path):
                # Search for any video file in temp_dir
                for f in os.listdir(temp_dir):
                    if f.endswith(('.mp4', '.mkv', '.webm', '.avi')):
                        video_path = os.path.join(temp_dir, f)
                        break
    except yt_dlp.utils.DownloadError as e:
        raise RuntimeError(f"下载失败，请检查链接是否有效。\n{e}")

    if not os.path.exists(video_path):
        raise RuntimeError("下载失败：未找到视频文件。")

    return video_path, title


def detect_scenes(video_path: str, threshold: float, progress_cb) -> list:
    """Detect scene cuts using PySceneDetect. Returns list of SceneCut."""
    progress_cb("正在分析镜头...", 0)

    video = open_video(video_path)
    scene_manager = SceneManager()
    scene_manager.add_detector(ContentDetector(threshold=threshold))

    total_frames = video.duration.get_frames()

    def frame_callback(frame_img, frame_num):
        if total_frames > 0 and frame_num % 30 == 0:
            pct = min((frame_num / total_frames) * 100, 100)
            progress_cb("正在分析镜头...", pct)

    scene_manager.detect_scenes(video, callback=frame_callback)
    scene_list = scene_manager.get_scene_list()

    if not scene_list:
        # No cuts detected — treat entire video as one scene
        video_cv = cv2.VideoCapture(video_path)
        fps = video_cv.get(cv2.CAP_PROP_FPS)
        frame_count = int(video_cv.get(cv2.CAP_PROP_FRAME_COUNT))
        duration = frame_count / fps if fps > 0 else 0
        video_cv.release()
        return [SceneCut(
            index=1,
            start_time=0.0,
            end_time=duration,
            start_frame=0,
            end_frame=frame_count - 1,
        )]

    scenes = []
    for i, (start_tc, end_tc) in enumerate(scene_list):
        scenes.append(SceneCut(
            index=i + 1,
            start_time=start_tc.get_seconds(),
            end_time=end_tc.get_seconds(),
            start_frame=start_tc.get_frames(),
            end_frame=end_tc.get_frames(),
        ))

    progress_cb("镜头分析完成", 100)
    return scenes


def capture_screenshots(video_path: str, scenes: list, output_dir: str,
                        progress_cb) -> list:
    """Capture a representative frame for each scene. Returns list of paths."""
    cap = cv2.VideoCapture(video_path)
    fps = cap.get(cv2.CAP_PROP_FPS)
    if fps <= 0:
        fps = 24.0
    offset_frames = int(fps * FRAME_OFFSET_SEC)

    shots_dir = os.path.join(output_dir, SHOT_DIR_NAME)
    os.makedirs(shots_dir, exist_ok=True)

    paths = []
    total = len(scenes)

    for i, scene in enumerate(scenes):
        target_frame = scene.start_frame + offset_frames
        if target_frame > scene.end_frame:
            target_frame = scene.start_frame

        cap.set(cv2.CAP_PROP_POS_FRAMES, target_frame)
        ret, frame = cap.read()

        filename = f"shot_{i + 1:03d}.jpg"
        filepath = os.path.join(shots_dir, filename)

        if ret and frame is not None:
            cv2.imwrite(filepath, frame, [cv2.IMWRITE_JPEG_QUALITY, 90])
            paths.append(filepath)
        else:
            paths.append(None)

        pct = ((i + 1) / total) * 100
        progress_cb(f"正在截取画面 {i + 1}/{total}...", pct)

    cap.release()
    return paths


def create_storyboard_excel(scenes: list, screenshot_paths: list,
                            output_path: str, progress_cb):
    """Generate the storyboard Excel file with embedded thumbnails."""
    progress_cb("正在生成分镜表...", 0)

    wb = Workbook()
    ws = wb.active
    ws.title = "分镜表"

    # Header
    headers = ["镜头编号", "时间码（起始）", "时长", "截图预览", "景别", "运镜", "备注"]
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    # Column widths
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 25
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["G"].width = 25

    ws.row_dimensions[1].height = 22

    total = len(scenes)
    center_align = Alignment(horizontal="center", vertical="center")

    for i, scene in enumerate(scenes):
        row = i + 2
        duration = scene.end_time - scene.start_time

        ws.cell(row=row, column=1, value=scene.index).alignment = center_align
        ws.cell(row=row, column=2, value=format_timecode(scene.start_time)).alignment = center_align
        ws.cell(row=row, column=3, value=round(duration, 2)).alignment = center_align
        # Columns E, F, G left blank for manual fill
        ws.cell(row=row, column=5, value="").alignment = center_align
        ws.cell(row=row, column=6, value="").alignment = center_align
        ws.cell(row=row, column=7, value="")

        # Embed thumbnail
        img_path = screenshot_paths[i] if i < len(screenshot_paths) else None
        if img_path and os.path.exists(img_path):
            try:
                pil_img = PILImage.open(img_path)
                orig_w, orig_h = pil_img.size
                pil_img.close()

                scale = THUMBNAIL_WIDTH_PX / orig_w if orig_w > 0 else 1
                new_h = int(orig_h * scale)

                xl_img = XlImage(img_path)
                xl_img.width = THUMBNAIL_WIDTH_PX
                xl_img.height = new_h
                xl_img.anchor = f"D{row}"
                ws.add_image(xl_img)

                # Row height in points (~ pixels * 0.75)
                ws.row_dimensions[row].height = max(new_h * 0.75, 20)
            except Exception:
                ws.row_dimensions[row].height = 20
        else:
            ws.row_dimensions[row].height = 20

        pct = ((i + 1) / total) * 100
        progress_cb(f"正在生成分镜表 {i + 1}/{total}...", pct)

    wb.save(output_path)
    progress_cb("分镜表生成完成", 100)


# ---------------------------------------------------------------------------
# Pipeline orchestrator
# ---------------------------------------------------------------------------

def run_pipeline(url: str, threshold: float, progress_cb, done_cb, error_cb):
    """Run the full analysis pipeline in a worker thread."""
    temp_dir = tempfile.mkdtemp(prefix="scene_storyboard_")

    try:
        # Phase 1: Download (0% - 40%)
        dl_cb = make_phase_cb(progress_cb, 0, 40)
        dl_cb("准备下载...", 0)
        video_path, title = download_video(url, temp_dir, dl_cb)

        # Phase 2: Scene detection (40% - 70%)
        detect_cb = make_phase_cb(progress_cb, 40, 70)
        scenes = detect_scenes(video_path, threshold, detect_cb)

        # Phase 3: Create output directory and capture screenshots (70% - 85%)
        output_dir = get_output_dir(title)
        capture_cb = make_phase_cb(progress_cb, 70, 85)
        screenshot_paths = capture_screenshots(video_path, scenes,
                                               str(output_dir), capture_cb)

        # Phase 4: Generate Excel (85% - 100%)
        excel_cb = make_phase_cb(progress_cb, 85, 100)
        excel_path = str(output_dir / EXCEL_FILENAME)
        create_storyboard_excel(scenes, screenshot_paths, excel_path, excel_cb)

        progress_cb("完成！", 100)
        done_cb(str(output_dir))

    except Exception as e:
        error_cb(str(e))
    finally:
        # Clean up temp directory
        try:
            shutil.rmtree(temp_dir, ignore_errors=True)
        except Exception:
            pass


# ---------------------------------------------------------------------------
# Tkinter UI
# ---------------------------------------------------------------------------

class App:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("视频镜头分镜工具")
        self.root.geometry("560x300")
        self.root.resizable(False, False)
        self._running = False
        self._build_widgets()

    def _build_widgets(self):
        pad = {"padx": 12, "pady": 4}

        # URL input
        tk.Label(self.root, text="视频链接：").grid(row=0, column=0, sticky="w", **pad)
        self.url_entry = tk.Entry(self.root, width=52)
        self.url_entry.grid(row=0, column=1, columnspan=2, sticky="we", **pad)

        # Threshold input
        tk.Label(self.root, text="检测阈值：").grid(row=1, column=0, sticky="w", **pad)
        self.threshold_entry = tk.Entry(self.root, width=10)
        self.threshold_entry.insert(0, str(DEFAULT_THRESHOLD))
        self.threshold_entry.grid(row=1, column=1, sticky="w", **pad)
        tk.Label(self.root, text="（值越小越灵敏）", fg="gray").grid(
            row=1, column=2, sticky="w")

        # Start button
        self.start_btn = tk.Button(self.root, text="开始分析", width=16,
                                   command=self._on_start)
        self.start_btn.grid(row=2, column=0, columnspan=3, pady=10)

        # Progress bar
        self.progress = ttk.Progressbar(self.root, length=500, mode="determinate")
        self.progress.grid(row=3, column=0, columnspan=3, **pad)

        # Status label
        self.status_label = tk.Label(self.root, text="就绪", fg="gray", anchor="w")
        self.status_label.grid(row=4, column=0, columnspan=3, sticky="we", **pad)

        # Output path (clickable)
        self.output_label = tk.Label(self.root, text="", fg="blue", cursor="hand2",
                                     anchor="w", wraplength=520)
        self.output_label.grid(row=5, column=0, columnspan=3, sticky="we", **pad)
        self.output_label.bind("<Button-1>", self._on_open_folder)
        self._output_path = None

        # Configure grid weight
        self.root.columnconfigure(1, weight=1)

    def _on_start(self):
        url = self.url_entry.get().strip()
        if not url:
            messagebox.showwarning("提示", "请输入视频链接")
            return

        try:
            threshold = float(self.threshold_entry.get().strip())
            if threshold <= 0:
                raise ValueError
        except ValueError:
            messagebox.showwarning("提示", "阈值请输入一个正数")
            return

        if self._running:
            return

        self._running = True
        self.start_btn.config(state="disabled")
        self.progress["value"] = 0
        self.status_label.config(text="正在启动...", fg="gray")
        self.output_label.config(text="")
        self._output_path = None

        thread = threading.Thread(
            target=run_pipeline,
            args=(url, threshold, self._on_progress, self._on_done, self._on_error),
            daemon=True,
        )
        thread.start()

    def _on_progress(self, message: str, percent: float):
        self.root.after(0, lambda: self._apply_progress(message, percent))

    def _apply_progress(self, message: str, percent: float):
        self.progress["value"] = min(percent, 100)
        self.status_label.config(text=message, fg="black")

    def _on_done(self, output_dir: str):
        self.root.after(0, lambda: self._apply_done(output_dir))

    def _apply_done(self, output_dir: str):
        self._running = False
        self.start_btn.config(state="normal")
        self.progress["value"] = 100
        self.status_label.config(text="分析完成！", fg="green")
        self._output_path = output_dir
        self.output_label.config(text=f"输出目录：{output_dir}（点击打开）")

    def _on_error(self, error_msg: str):
        self.root.after(0, lambda: self._apply_error(error_msg))

    def _apply_error(self, error_msg: str):
        self._running = False
        self.start_btn.config(state="normal")
        self.progress["value"] = 0
        self.status_label.config(text="出错了", fg="red")
        messagebox.showerror("错误", error_msg)

    def _on_open_folder(self, event=None):
        if self._output_path and os.path.isdir(self._output_path):
            try:
                open_folder(self._output_path)
            except Exception:
                pass


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def main():
    root = tk.Tk()
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
